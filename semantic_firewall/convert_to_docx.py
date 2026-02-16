"""Convert TECHNICAL_ARCHITECTURE.md to a formatted DOCX document."""

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def parse_md(md_text: str):
    """Parse markdown into structured blocks."""
    blocks = []
    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]

        # Headings
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            blocks.append(("heading", level, text))
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ("---", "***", "___"):
            blocks.append(("hr",))
            i += 1
            continue

        # Table
        if "|" in line and i + 1 < len(lines) and re.match(r"^\s*\|[\s\-:|]+\|", lines[i + 1]):
            table_lines = []
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i])
                i += 1
            # Parse table
            rows = []
            for tl in table_lines:
                if re.match(r"^\s*\|[\s\-:|]+\|$", tl):
                    continue  # separator row
                cells = [c.strip() for c in tl.split("|")]
                cells = [c for c in cells if c != ""]
                if cells:
                    rows.append(cells)
            if rows:
                blocks.append(("table", rows))
            continue

        # Code block
        if line.strip().startswith("```"):
            lang = line.strip().lstrip("`").strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            blocks.append(("code", "\n".join(code_lines), lang))
            continue

        # Bullet list
        if re.match(r"^\s*[-*]\s", line):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s", lines[i]):
                indent = len(lines[i]) - len(lines[i].lstrip())
                text = re.sub(r"^\s*[-*]\s+", "", lines[i])
                items.append((indent, text))
                i += 1
            blocks.append(("bullet_list", items))
            continue

        # Numbered list
        if re.match(r"^\s*\d+\.\s", line):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s", lines[i]):
                text = re.sub(r"^\s*\d+\.\s+", "", lines[i])
                items.append(text)
                i += 1
            blocks.append(("numbered_list", items))
            continue

        # Regular paragraph
        if line.strip():
            para_lines = []
            while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") and not lines[i].strip().startswith("```") and not lines[i].strip().startswith("|") and not lines[i].strip() in ("---", "***", "___"):
                para_lines.append(lines[i])
                i += 1
            blocks.append(("paragraph", " ".join(para_lines)))
            continue

        i += 1

    return blocks


def add_formatted_text(paragraph, text):
    """Add text with inline markdown formatting (bold, code, italic)."""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def set_cell_shading(cell, color):
    """Set cell background color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shading = tcPr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): color,
    })
    tcPr.append(shading)


def build_docx(blocks, output_path):
    """Build DOCX from parsed blocks."""
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Configure heading styles
    for level in range(1, 5):
        hs = doc.styles[f"Heading {level}"]
        hs.font.name = "Calibri"
        hs.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        if level == 1:
            hs.font.size = Pt(24)
        elif level == 2:
            hs.font.size = Pt(18)
        elif level == 3:
            hs.font.size = Pt(14)
        else:
            hs.font.size = Pt(12)

    for block in blocks:
        btype = block[0]

        if btype == "heading":
            level = min(block[1], 4)
            text = block[2]
            # Skip the table of contents links
            text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
            h = doc.add_heading(text, level=level)

        elif btype == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("_" * 80)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(8)

        elif btype == "paragraph":
            text = block[1]
            # Convert markdown links
            text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_formatted_text(p, text)

        elif btype == "code":
            code = block[1]
            lang = block[2] if len(block) > 2 else ""
            # Add language label if present
            if lang:
                lp = doc.add_paragraph()
                run = lp.add_run(f"  {lang}")
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
                run.font.name = "Consolas"
                lp.paragraph_format.space_after = Pt(0)

            for code_line in code.split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Inches(0.3)
                run = p.add_run(code_line)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

        elif btype == "table":
            rows = block[1]
            if not rows:
                continue
            num_cols = max(len(r) for r in rows)

            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT

            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    if ci < num_cols:
                        cell = table.cell(ri, ci)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        # Clean markdown from cell text
                        clean = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", cell_text)
                        add_formatted_text(p, clean)
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)

                        for run in p.runs:
                            run.font.size = Pt(9)

                        # Header row styling
                        if ri == 0:
                            set_cell_shading(cell, "1E293B")
                            for run in p.runs:
                                run.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                run.font.size = Pt(9)

            doc.add_paragraph()  # spacing after table

        elif btype == "bullet_list":
            items = block[1]
            for indent, text in items:
                text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
                level = min(indent // 2, 3)
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
                p.paragraph_format.space_after = Pt(2)
                add_formatted_text(p, text)

        elif btype == "numbered_list":
            items = block[1]
            for text in items:
                text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(2)
                add_formatted_text(p, text)

    doc.save(str(output_path))
    print(f"Saved: {output_path}")


def main():
    md_path = Path(__file__).parent / "TECHNICAL_ARCHITECTURE.md"
    docx_path = Path(__file__).parent / "TECHNICAL_ARCHITECTURE.docx"

    md_text = md_path.read_text(encoding="utf-8")
    blocks = parse_md(md_text)
    build_docx(blocks, docx_path)


if __name__ == "__main__":
    main()
